"""Documentos de datos tabulares con texto escapado y encabezados repetidos."""
from datetime import date,datetime
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4,landscape
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,TableStyle
from configuracion import TZ

def texto(v):
    if v is None:return '—'
    if isinstance(v,datetime):return v.astimezone(TZ).strftime('%d/%m/%Y %H:%M') if v.tzinfo else v.strftime('%d/%m/%Y %H:%M')
    if isinstance(v,date):return v.strftime('%d/%m/%Y')
    if isinstance(v,Decimal):return f'{v:,.2f}'
    return str(v)

def exportar(reporte,formato):
    salida=BytesIO();columnas=reporte['columnas']
    periodo='Corte actual' if reporte['tipo']=='inventario' else f"{reporte['desde']} al {reporte['hasta']}"
    resumen=f"Registros: {reporte['cantidad']}"+('' if reporte['tipo']=='inventario' else f" | Neto/total Q: {texto(reporte['total'])}")
    if formato=='excel':
        libro=Workbook();hoja=libro.active;hoja.title='Punto Com'
        for fila in ([reporte['titulo']],[periodo],[reporte['nota']],[f"Generado: {texto(reporte['generado'])} - Guatemala"],
                     [reporte['responsable']], [resumen],[],[c[1] for c in columnas]):
            hoja.append(fila)
        for fila in reporte['data']:hoja.append([fila.get(k) for k,_ in columnas])
        if not reporte['data']:hoja.append(['Sin registros para la selección.'])
        for fila in hoja:
            for celda in fila:
                if isinstance(celda.value,str):celda.data_type='s'
                if isinstance(celda.value,datetime):
                    if celda.value.tzinfo:celda.value=celda.value.astimezone(TZ).replace(tzinfo=None)
                    celda.number_format='DD/MM/YYYY HH:MM'
                if isinstance(celda.value,Decimal):celda.number_format='#,##0.00'
        for c in hoja[8]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='173456')
        for i in range(1,len(columnas)+1):hoja.column_dimensions[hoja.cell(8,i).column_letter].width=28
        hoja.freeze_panes='A9';hoja.auto_filter.ref=f'A8:{hoja.cell(max(8,hoja.max_row),len(columnas)).coordinate}'
        libro.save(salida)
    else:
        estilos=getSampleStyleSheet();celda=ParagraphStyle('celda',fontName='Helvetica',fontSize=8,leading=11,wordWrap='CJK')
        def par(v):return Paragraph(escape(texto(v)).replace('\n','<br/>'),celda)
        doc=SimpleDocTemplate(salida,pagesize=landscape(A4),rightMargin=30,leftMargin=30,topMargin=30,bottomMargin=34)
        contenido=[Paragraph('PUNTO COM',estilos['Title']),Paragraph(escape(reporte['titulo']),estilos['Heading2']),
          par(periodo),par(reporte['nota']),par('Generado: '+texto(reporte['generado'])+' - Guatemala | '+reporte['responsable']),
          par(resumen),Spacer(1,14)]
        tabla=LongTable([[par(c[1]) for c in columnas]]+[[par(fila.get(k)) for k,_ in columnas] for fila in reporte['data']],
          colWidths=[(landscape(A4)[0]-60)/len(columnas)]*len(columnas),repeatRows=1,hAlign='LEFT')
        tabla.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e6edf7')),('VALIGN',(0,0),(-1,-1),'TOP'),
          ('GRID',(0,0),(-1,-1),.3,colors.HexColor('#dce3ed')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f6f8fb')]),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
        contenido.append(tabla)
        if not reporte['data']:contenido.append(par('Sin registros para la selección.'))
        def pie(canvas,documento):
            canvas.setFont('Helvetica',8);canvas.drawRightString(landscape(A4)[0]-30,18,'Punto Com | Página '+str(documento.page))
        doc.build(contenido,onFirstPage=pie,onLaterPages=pie)
    salida.seek(0);return salida
