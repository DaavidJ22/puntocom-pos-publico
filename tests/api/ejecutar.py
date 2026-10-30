"""Pruebas JWT y API sobre una base desechable; no escribe en la base de trabajo."""

from datetime import timedelta
import json
from pathlib import Path
import sys
import unittest

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
sys.path.insert(0,str(ROOT/'tests'/'postgres'))
from sqlalchemy import text
from flask_jwt_extended import create_access_token, decode_token
from entorno_pg import BaseTemporalPostgres,CasoPostgres,configuracion_origen
from backend.app import create_app
from datos import db


class APIJWTTests(CasoPostgres):
    aplicacion=None

    def setUp(self):
        super().setUp()
        self.client=self.aplicacion.test_client()

    def tearDown(self):
        with self.aplicacion.app_context():
            db.session.remove()
        super().tearDown()

    def login(self,nombre='admin',client=None):
        client=client or self.client
        respuesta=client.post('/api/auth/login',json={'correo':nombre+'@example.test','password':'PruebaPostgres2026!'})
        self.assertEqual(respuesta.status_code,200,respuesta.json)
        return respuesta.json['access_token']

    def headers(self,token):
        return {'Authorization':'Bearer '+token}

    def csrf(self):
        return {'X-CSRF-TOKEN':self.client.get_cookie('csrf_refresh_token').value}

    def post(self,path,token,body):
        return self.client.post('/api'+path,headers=self.headers(token),json=body)

    def test_health_comprueba_conexion_real_y_responde_503_si_falla(self):
        from unittest.mock import patch
        from sqlalchemy.exc import OperationalError
        r=self.client.get('/api/health')
        self.assertEqual(r.status_code,200)
        self.assertEqual(r.json['base_datos'],'conectada')
        with patch('backend.routes.value',side_effect=OperationalError('SELECT 1',{},Exception('sin conexión'))):
            r=self.client.get('/api/health')
        self.assertEqual(r.status_code,503)
        self.assertEqual(r.json['base_datos'],'no disponible')

    def test_login_cookie_httponly_sin_hashes_y_perfiles(self):
        for nombre,rol in [('admin','Administrador'),('cajero1','Cajero'),('bodeguero','Bodeguero')]:
            token=self.login(nombre)
            me=self.client.get('/api/auth/me',headers=self.headers(token))
            self.assertEqual(me.json['rol'],rol)
            self.assertNotIn('hash_contrasena',me.json)
            with self.aplicacion.app_context():
                claims=decode_token(token)
            self.assertEqual(claims['type'],'access')
            self.assertEqual(claims['exp']-claims['iat'],600)
        cookie=self.client.get_cookie('refresh_token_cookie',path='/api/auth')
        self.assertTrue(cookie.http_only)
        self.assertEqual(cookie.same_site,'Strict')

    def test_credenciales_invalidas_no_crean_sesion(self):
        r=self.client.post('/api/auth/login',json={'correo':'admin@example.test','password':'incorrecta'})
        self.assertEqual(r.status_code,401)
        with self.engine.connect() as c:
            self.assertEqual(c.execute(text('SELECT COUNT(*) FROM sesion_token')).scalar(),0)

    def test_sin_token_y_token_manipulado_rechazados(self):
        for headers in ({},{'Authorization':'Bearer token-falso'}):
            self.assertEqual(self.client.get('/api/registros/productos',headers=headers).status_code,401)

    def test_access_vencido_puede_renovarse_con_refresh(self):
        token=self.login()
        with self.aplicacion.app_context():
            claims=decode_token(token)
            vencido=create_access_token(identity=claims['sub'],additional_claims={'sid':claims['sid']},expires_delta=timedelta(seconds=-1))
        self.assertEqual(self.client.get('/api/auth/me',headers=self.headers(vencido)).status_code,401)
        r=self.client.post('/api/auth/refresh',headers=self.csrf())
        self.assertEqual(r.status_code,200)
        self.assertEqual(self.client.get('/api/auth/me',headers=self.headers(r.json['access_token'])).status_code,200)

    def test_refresh_exige_csrf(self):
        self.login()
        self.assertEqual(self.client.post('/api/auth/refresh').status_code,401)

    def test_access_no_sustituye_refresh(self):
        token=self.login()
        nuevo=self.aplicacion.test_client()
        self.assertEqual(nuevo.post('/api/auth/refresh',headers=self.headers(token)).status_code,401)

    def test_rotacion_y_reutilizacion_revocan_familia(self):
        token=self.login()
        old=self.client.get_cookie('refresh_token_cookie',path='/api/auth').value
        csrf=self.client.get_cookie('csrf_refresh_token').value
        r=self.client.post('/api/auth/refresh',headers=self.csrf())
        self.assertEqual(r.status_code,200,r.json)
        self.assertNotEqual(old,self.client.get_cookie('refresh_token_cookie',path='/api/auth').value)
        new_token=r.json['access_token']
        self.client.set_cookie('refresh_token_cookie',old,path='/api/auth')
        self.client.set_cookie('csrf_refresh_token',csrf)
        r=self.client.post('/api/auth/refresh',headers=self.csrf())
        self.assertEqual(r.status_code,401)
        self.assertEqual(r.json['code'],'refresh_reused')
        for access in (token,new_token):
            self.assertEqual(self.client.get('/api/auth/me',headers=self.headers(access)).status_code,401)

    def test_logout_revoca_access_y_refresh_persistentemente(self):
        token=self.login()
        cookie=self.client.get_cookie('refresh_token_cookie',path='/api/auth').value
        csrf=self.client.get_cookie('csrf_refresh_token').value
        self.assertEqual(self.client.post('/api/auth/logout',headers=self.csrf()).status_code,200)
        self.assertEqual(self.client.get('/api/auth/me',headers=self.headers(token)).status_code,401)
        self.client.set_cookie('refresh_token_cookie',cookie,path='/api/auth')
        self.client.set_cookie('csrf_refresh_token',csrf)
        self.assertEqual(self.client.post('/api/auth/refresh',headers=self.csrf()).status_code,401)

    def test_sesiones_distintas_no_se_revocan_juntas(self):
        primero=self.login()
        otra=self.aplicacion.test_client();segundo=self.login(client=otra)
        self.client.post('/api/auth/logout',headers=self.csrf())
        self.assertEqual(otra.get('/api/auth/me',headers=self.headers(segundo)).status_code,200)
        self.assertEqual(self.client.get('/api/auth/me',headers=self.headers(primero)).status_code,401)

    def test_desactivacion_inmediata(self):
        token=self.login('cajero1')
        with self.engine.begin() as c:c.execute(text('UPDATE usuario SET estado=FALSE WHERE id_usuario=2'))
        self.assertEqual(self.client.get('/api/auth/me',headers=self.headers(token)).status_code,401)
        self.assertEqual(self.client.post('/api/auth/refresh',headers=self.csrf()).status_code,401)

    def test_cambio_rol_inmediato_sin_confiar_en_claims(self):
        token=self.login()
        with self.engine.begin() as c:c.execute(text('UPDATE usuario SET id_rol=3 WHERE id_usuario=1'))
        self.assertEqual(self.client.get('/api/dashboard',headers=self.headers(token)).status_code,403)
        self.assertEqual(self.client.get('/api/auth/me',headers=self.headers(token)).json['rol'],'Bodeguero')

    def test_cajero_no_ve_costos_ni_administra(self):
        token=self.login('cajero1')
        r=self.client.get('/api/registros/productos',headers=self.headers(token))
        self.assertEqual(r.status_code,200)
        self.assertNotIn('costo',r.json['data'][0])
        self.assertEqual(self.post('/registros/productos',token,{}).status_code,403)
        self.assertEqual(self.client.get('/api/registros/usuarios',headers=self.headers(token)).status_code,403)
        self.assertEqual(self.client.get('/api/inventario',headers=self.headers(token)).status_code,403)

    def test_bodeguero_no_cobra_ni_consulta_ventas(self):
        token=self.login('bodeguero')
        self.assertEqual(self.post('/ventas',token,{}).status_code,403)
        self.assertEqual(self.client.get('/api/ventas',headers=self.headers(token)).status_code,403)
        self.assertEqual(self.client.get('/api/inventario',headers=self.headers(token)).status_code,200)

    def payload(self,**cambios):
        return {'cliente_id':1,'metodo':'EFECTIVO','recibido':'100','items':[{'id':1,'cantidad':1,'descuento':0}],**cambios}

    def test_cobro_recalcula_precio_y_ticket_es_privado(self):
        token=self.login('cajero1')
        d=self.payload();d['items'][0]['precio_unitario']=0.01
        r=self.post('/ventas',token,d)
        self.assertEqual(r.status_code,403,r.json)
        r=self.post('/ventas',token,self.payload())
        self.assertEqual(r.status_code,200,r.json)
        vid=r.json['resultado']['id_venta']
        ticket=self.client.get('/api/ventas/'+str(vid),headers=self.headers(token))
        self.assertEqual(ticket.json['venta']['total'],'15.00')
        otro=self.login('cajero2')
        self.assertEqual(self.client.get('/api/ventas/'+str(vid),headers=self.headers(otro)).status_code,403)
        self.assertEqual(self.client.get('/api/reportes',headers=self.headers(otro)).json['ventas'],[])

    def test_cajero_no_inyecta_descuentos(self):
        token=self.login('cajero1');d=self.payload();d['items'][0]['descuento']=1
        self.assertEqual(self.post('/ventas',token,d).status_code,403)

    def test_cobro_sin_stock_revierte_todo(self):
        token=self.login('cajero1');d=self.payload();d['items'][0]['cantidad']=99
        self.assertEqual(self.post('/ventas',token,d).status_code,400)
        with self.engine.connect() as c:
            self.assertEqual(c.execute(text('SELECT COUNT(*) FROM venta')).scalar(),0)
            self.assertEqual(c.execute(text('SELECT existencia_actual FROM vista_inventario WHERE id_producto=1')).scalar(),10)

    def test_caja_admin_cobro_anulacion_y_excel(self):
        token=self.login()
        self.assertEqual(self.post('/caja',token,{'accion':'abrir','fondo_inicial':'100'}).status_code,200)
        r=self.post('/ventas',token,self.payload());self.assertEqual(r.status_code,200,r.json)
        vid=r.json['resultado']['id_venta']
        self.assertEqual(self.client.get('/api/reportes/excel',headers=self.headers(token)).status_code,200)
        self.assertEqual(self.post('/ventas/'+str(vid)+'/anular',token,{'motivo':'Error de prueba'}).status_code,200)
        self.assertEqual(self.client.get('/api/ventas/'+str(vid),headers=self.headers(token)).json['venta']['estado'],'ANULADA')

    def test_cliente_duplicado_y_validacion(self):
        token=self.login('cajero1');d={'nombre_razon_social':'Ana Prueba','nit_dpi':'1234567K','telefono':'12345678'}
        self.assertEqual(self.post('/registros/clientes',token,d).status_code,200)
        self.assertEqual(self.post('/registros/clientes',token,d).status_code,400)
        self.assertEqual(self.post('/registros/clientes',token,{'nombre_razon_social':''}).status_code,400)

    def test_conteo_aprobado_no_duplica_ajuste(self):
        token=self.login('bodeguero')
        r=self.post('/inventario/conteos',token,{'id_producto':1,'cantidad_contada':12,'motivo':'Conteo de prueba'})
        self.assertEqual(r.status_code,200,r.json);aid=r.json['resultado']
        admin=self.login()
        self.assertEqual(self.post('/inventario/ajustes/'+str(aid)+'/aprobar',admin,{}).status_code,200)
        self.assertEqual(self.post('/inventario/ajustes/'+str(aid)+'/aprobar',admin,{}).status_code,400)
        with self.engine.connect() as c:self.assertEqual(c.execute(text('SELECT existencia_actual FROM vista_inventario WHERE id_producto=1')).scalar(),12)

    def test_origen_externo_rechazado(self):
        self.assertEqual(self.client.post('/api/auth/login',headers={'Origin':'https://otro.invalid'},json={}).status_code,403)

    def test_origen_lan_explicito_no_autoriza_otros_origenes(self):
        from unittest.mock import patch
        with patch.dict('os.environ',{'FRONTEND_ORIGINS':'https://192.168.1.136:4173'}):
            app=create_app(self.aplicacion.config)
        try:
            cliente=app.test_client()
            self.assertEqual(cliente.post('/api/auth/login',headers={'Origin':'https://192.168.1.136:4173'},json={}).status_code,401)
            self.assertEqual(cliente.post('/api/auth/login',headers={'Origin':'https://192.168.1.137:4173'},json={}).status_code,403)
        finally:
            with app.app_context():db.session.remove();db.engine.dispose()

    def test_produccion_https_cookie_secure_refresh_y_logout(self):
        from unittest.mock import patch
        origen='https://puntocom-prueba.vercel.app'
        destino='https://puntocom-api-prueba.onrender.com'
        with patch.dict('os.environ',{'FRONTEND_ORIGIN':origen,'FRONTEND_ORIGINS':''}):
            app=create_app({**self.aplicacion.config,'APP_ENV':'production','SESSION_COOKIE_SECURE':True,'TRUST_PROXY':True})
        try:
            cliente=app.test_client()
            r=cliente.post('/api/auth/login',base_url=destino,headers={'Origin':origen},json={'correo':'admin@example.test','password':'PruebaPostgres2026!'})
            self.assertEqual(r.status_code,200,r.json)
            cookie=cliente.get_cookie('refresh_token_cookie',domain='puntocom-api-prueba.onrender.com',path='/api/auth')
            self.assertTrue(cookie.secure);self.assertTrue(cookie.http_only)
            self.assertEqual(cookie.same_site,'Strict')
            csrf=cliente.get_cookie('csrf_refresh_token',domain='puntocom-api-prueba.onrender.com').value
            r=cliente.post('/api/auth/refresh',base_url=destino,headers={'Origin':origen,'X-CSRF-TOKEN':csrf})
            self.assertEqual(r.status_code,200,r.json)
            self.assertEqual(r.headers['Cache-Control'],'no-store')
            self.assertEqual(r.headers['CDN-Cache-Control'],'no-store')
            csrf=cliente.get_cookie('csrf_refresh_token',domain='puntocom-api-prueba.onrender.com').value
            self.assertEqual(cliente.post('/api/auth/logout',base_url=destino,headers={'Origin':origen,'X-CSRF-TOKEN':csrf}).status_code,200)
            self.assertIsNone(cliente.get_cookie('refresh_token_cookie',domain='puntocom-api-prueba.onrender.com',path='/api/auth'))
        finally:
            with app.app_context():db.session.remove();db.engine.dispose()

    def test_json_invalido_y_modulo_desconocido(self):
        token=self.login()
        self.assertEqual(self.client.post('/api/registros/clientes',headers=self.headers(token),json=[]).status_code,400)
        self.assertEqual(self.client.get('/api/registros/no_existe',headers=self.headers(token)).status_code,404)

    def test_filtros_fechas_invalidas_no_exportan(self):
        token=self.login()
        self.assertEqual(self.client.get('/api/reportes?desde=2026-10-20&hasta=2026-10-01',headers=self.headers(token)).status_code,400)

    def test_producto_nuevo_edicion_y_proteccion_de_historial(self):
        token=self.login()
        d={'id_categoria':1,'codigo':'API001','nombre':'Producto de prueba','marca':'Marca',
           'unidad':'Unidad','costo':'5','precio_venta':'8','stock_minimo':2,'stock_inicial':3}
        r=self.post('/registros/productos',token,d)
        self.assertEqual(r.status_code,200,r.json);pid=r.json['resultado']
        r=self.client.put('/api/registros/productos/'+str(pid),headers=self.headers(token),json={**d,'precio_venta':'9'})
        self.assertEqual(r.status_code,200,r.json)
        listado=self.client.get('/api/registros/productos',headers=self.headers(token)).json['data']
        p=next(p for p in listado if p['id_producto']==pid)
        self.assertEqual(p['precio_venta'],'9.00');self.assertEqual(p['existencia_actual'],3)
        self.assertFalse(p['puede_eliminar'])
        self.assertEqual(self.client.delete('/api/registros/productos/'+str(pid),headers=self.headers(token)).status_code,400)

    def test_compra_actualiza_stock_y_costo(self):
        token=self.login()
        r=self.post('/compras',token,{'id_producto':1,'id_proveedor':1,'cantidad':2,'costo_unitario':'7.00','metodo_pago':'TRANSFERENCIA','numero_factura':'API-PRUEBA'})
        self.assertEqual(r.status_code,200,r.json)
        with self.engine.connect() as c:
            self.assertEqual(c.execute(text('SELECT existencia_actual FROM vista_inventario WHERE id_producto=1')).scalar(),12)
            self.assertEqual(str(c.execute(text('SELECT costo FROM producto WHERE id_producto=1')).scalar()),'7.00')

    def test_flujo_completo_con_relaciones_y_persistencia(self):
        from scripts.auditar_bd import auditar
        token=self.login()
        r=self.post('/compras',token,{'id_producto':1,'id_proveedor':1,'cantidad':2,'costo_unitario':'7.00','metodo_pago':'TRANSFERENCIA','numero_factura':'FLUJO'})
        self.assertEqual(r.status_code,200,r.json)
        self.assertEqual(self.post('/caja',token,{'accion':'abrir','fondo_inicial':'100.00'}).status_code,200)
        r=self.post('/ventas',token,{'cliente_id':1,'metodo':'EFECTIVO','recibido':'20.00','items':[{'id':1,'cantidad':1}]})
        self.assertEqual(r.status_code,200,r.json)
        vid=r.json['resultado']['id_venta']
        # Otra petición y otra conexión deben observar lo confirmado.
        self.assertEqual(self.client.get('/api/ventas/'+str(vid),headers=self.headers(token)).json['detalle'][0]['cantidad'],1)
        self.assertEqual(self.post('/ventas/'+str(vid)+'/anular',token,{'motivo':'Comprobación de relaciones'}).status_code,200)
        self.assertEqual(self.post('/caja',token,{'accion':'cerrar','efectivo_contado':'100.00'}).status_code,200)
        with self.engine.connect() as c:
            compra=c.execute(text('''SELECT m.cantidad,m.id_producto=d.id_producto AS coincide
                FROM movimiento_inventario m JOIN detalle_compra d USING(id_detalle_compra)''')).mappings().one()
            self.assertEqual(compra['cantidad'],2);self.assertTrue(compra['coincide'])
            inventario=c.execute(text('''SELECT m.tipo,m.cantidad,m.id_producto=d.id_producto AS coincide
                FROM movimiento_inventario m JOIN detalle_venta d USING(id_detalle_venta)
                WHERE d.id_venta=:id ORDER BY m.tipo'''),{'id':vid}).mappings().all()
            self.assertEqual([r['tipo'] for r in inventario],['ENTRADA','SALIDA'])
            self.assertTrue(all(r['coincide'] and r['cantidad']==1 for r in inventario))
            movimientos=c.execute(text('SELECT tipo,monto FROM movimiento_caja WHERE id_venta=:id ORDER BY tipo'),{'id':vid}).all()
            self.assertEqual([(tipo,str(monto)) for tipo,monto in movimientos],[('EGRESO','15.00'),('INGRESO','15.00')])
            self.assertEqual(c.execute(text('SELECT existencia_actual FROM vista_inventario WHERE id_producto=1')).scalar_one(),12)
        self.assertTrue(auditar(self.engine)['controles_sin_incidencias'])

    def test_alta_usuario_y_cambio_rol_protegen_ultimo_admin(self):
        token=self.login()
        r=self.client.put('/api/registros/usuarios/1',headers=self.headers(token),json={'id_rol':2})
        self.assertEqual(r.status_code,400)
        r=self.post('/registros/usuarios',token,{'nombre':'Usuario de prueba','correo':'api@example.test','password':'ClaveLarga2026!','id_rol':2})
        self.assertEqual(r.status_code,200,r.json)
        nuevo=self.client.post('/api/auth/login',json={'correo':'api@example.test','password':'ClaveLarga2026!'})
        self.assertEqual(nuevo.status_code,200)
        self.assertEqual(self.client.get('/api/dashboard',headers=self.headers(nuevo.json['access_token'])).status_code,403)

    def test_inventario_excel_no_filtra_costos_al_bodeguero(self):
        from openpyxl import load_workbook
        from io import BytesIO
        token=self.login('bodeguero')
        r=self.client.get('/api/inventario/excel',headers=self.headers(token))
        self.assertEqual(r.status_code,200)
        libro=load_workbook(BytesIO(r.data));cabeceras=[c.value for c in libro.active[1]];libro.close()
        self.assertNotIn('costo',cabeceras)

    def eventos(self,token,query=''):
        r=self.client.get('/api/auditoria'+query,headers=self.headers(token))
        self.assertEqual(r.status_code,200,r.json)
        return r.json

    def test_bitacora_solo_administrador_incluso_por_peticion_directa(self):
        self.assertEqual(self.client.get('/api/auditoria').status_code,401)
        for nombre in ('cajero1','bodeguero'):
            token=self.login(nombre)
            self.assertEqual(self.client.get('/api/auditoria',headers=self.headers(token)).status_code,403)
        token=self.login()
        self.assertIn('auditoria.ver',self.client.get('/api/auth/me',headers=self.headers(token)).json['permisos'])
        self.assertEqual(self.eventos(token)['pagina'],1)
        for method in ('post','put','delete'):
            self.assertEqual(getattr(self.client,method)('/api/auditoria',headers=self.headers(token)).status_code,405)

    def test_bitacora_precio_antes_despues_y_actor_no_inyectable(self):
        token=self.login()
        d={'id_categoria':1,'codigo':'PG001','nombre':'Cuaderno Uno','unidad':'Unidad',
           'costo':'9','precio_venta':'18.50','stock_minimo':5,'id_usuario':999,'usuario_nombre':'Falso'}
        r=self.client.put('/api/registros/productos/1',headers=self.headers(token),json=d)
        self.assertEqual(r.status_code,200,r.json)
        evento=self.eventos(token,'?modulo=productos&accion=ACTUALIZAR')['data'][0]
        self.assertEqual(evento['antes']['precio_venta'],15)
        self.assertEqual(evento['despues']['precio_venta'],18.5)
        self.assertEqual(evento['id_usuario'],1)
        self.assertEqual(evento['usuario_nombre'],'Administrador')
        self.assertEqual(evento['origen'],'api')
        self.assertIsNotNone(evento['solicitud'])

    def test_bitacora_rollback_de_compra_no_deja_historial_falso(self):
        from unittest.mock import patch
        import backend.routes as rutas
        token=self.login()
        with self.engine.connect() as c:
            antes=c.execute(text('SELECT COUNT(*) FROM bitacora_auditoria')).scalar_one()
        original=rutas.value
        def falla_despues_de_insertar_compra(sql,**params):
            if 'INSERT INTO detalle_compra' in sql:
                raise ValueError('Fallo controlado después del encabezado')
            return original(sql,**params)
        with patch('backend.routes.value',side_effect=falla_despues_de_insertar_compra):
            r=self.post('/compras',token,{'id_producto':1,'id_proveedor':1,'cantidad':2,'costo_unitario':'7',
                                      'metodo_pago':'EFECTIVO','numero_factura':'ROLLBACK'})
        self.assertEqual(r.status_code,400,r.json)
        with self.engine.connect() as c:
            self.assertEqual(c.execute(text('SELECT COUNT(*) FROM compra')).scalar_one(),0)
            self.assertEqual(c.execute(text('SELECT COUNT(*) FROM bitacora_auditoria')).scalar_one(),antes)

    def test_bitacora_solicitud_relaciona_venta_inventario_caja_y_anulacion(self):
        token=self.login('cajero1')
        r=self.post('/ventas',token,self.payload());self.assertEqual(r.status_code,200,r.json)
        vid=r.json['resultado']['id_venta']
        admin=self.login()
        eventos=self.eventos(admin,'?modulo=ventas&accion=CREAR&registro='+str(vid))['data']
        venta=next(e for e in eventos if e['entidad']=='venta')
        with self.engine.connect() as c:
            filas=c.execute(text('SELECT entidad,id_usuario FROM bitacora_auditoria WHERE solicitud=CAST(:s AS uuid)'),{'s':venta['solicitud']}).all()
        self.assertTrue({'venta','detalle_venta','movimiento_inventario','movimiento_caja'} <= {e[0] for e in filas})
        self.assertTrue(all(e[1]==2 for e in filas))
        self.assertEqual(self.post('/caja',admin,{'accion':'abrir','fondo_inicial':'100'}).status_code,200)
        self.assertEqual(self.post('/ventas/'+str(vid)+'/anular',admin,{'motivo':'Error de digitación'}).status_code,200)
        evento=self.eventos(admin,'?accion=ANULAR')['data'][0]
        self.assertEqual(evento['id_usuario'],1)
        self.assertEqual(evento['despues']['motivo_anulacion'],'Error de digitación')

    def test_bitacora_accesos_sin_credenciales_y_sin_suplantar_al_usuario_rechazado(self):
        clave='NoSeDebeGuardarEstaClave!'
        self.client.post('/api/auth/login',json={'correo':'admin@example.test','password':clave})
        token=self.login()
        self.assertEqual(self.client.post('/api/auth/refresh',headers=self.csrf()).status_code,200)
        self.assertEqual(self.client.post('/api/auth/logout',headers=self.csrf()).status_code,200)
        with self.engine.connect() as c:
            filas=[dict(r) for r in c.execute(text("SELECT * FROM bitacora_auditoria WHERE modulo='accesos'")).mappings()]
        acciones={r['accion'] for r in filas}
        self.assertTrue({'INICIAR_SESION','INGRESO_RECHAZADO','RENOVAR_SESION','CERRAR_SESION'}<=acciones)
        self.assertIsNone(next(r for r in filas if r['accion']=='INGRESO_RECHAZADO')['id_usuario'])
        for privado in (clave,token,'hash_contrasena','refresh_hash','csrf_refresh_token'):
            self.assertNotIn(privado,str(filas))

    def test_bitacora_alta_usuario_no_expone_hash_o_clave(self):
        token=self.login()
        clave='SoloParaCrearUsuario2026!'
        r=self.post('/registros/usuarios',token,{'nombre':'Usuario de prueba','correo':'audit@example.test','password':clave,'id_rol':2})
        self.assertEqual(r.status_code,200,r.json)
        eventos=self.eventos(token,'?modulo=usuarios')['data']
        self.assertNotIn(clave,str(eventos));self.assertNotIn('hash_contrasena',str(eventos))
        self.assertEqual(eventos[0]['despues']['correo'],'audit@example.test')

    def test_bitacora_filtros_paginacion_validacion_y_fechas_de_guatemala(self):
        from configuracion import TZ
        from datetime import datetime
        token=self.login()
        for sufijo in ('Uno','Dos','Tres'):
            self.assertEqual(self.post('/registros/categorias',token,{'nombre':'Auditoria '+sufijo}).status_code,200)
        from urllib.parse import urlencode
        hoy=datetime.now(TZ).date().isoformat()
        query='?'+urlencode(dict(modulo='categorias',accion='CREAR',usuario=1,desde=hoy,hasta=hoy,limite=2))
        primero=self.eventos(token,query)
        segundo=self.eventos(token,query+'&pagina=2')
        self.assertEqual(primero['total'],3);self.assertEqual(len(primero['data']),2)
        self.assertEqual(len(segundo['data']),1)
        self.assertTrue({r['id_evento'] for r in primero['data']}.isdisjoint(r['id_evento'] for r in segundo['data']))
        self.assertEqual(self.eventos(token,'?registro=99999999')['data'],[])
        for filtro in ('pagina=0','limite=101','modulo=no_existe','usuario=x',"accion=' OR 1=1 --",'desde=2026-10-20&hasta=2026-10-01','registro='+'x'*101):
            r=self.client.get('/api/auditoria?'+filtro,headers=self.headers(token))
            self.assertEqual(r.status_code,400,r.json)

    def test_bitacora_api_no_conserva_actor_entre_solicitudes(self):
        admin=self.login()
        cajero=self.login('cajero1')
        self.assertEqual(self.post('/registros/categorias',admin,{'nombre':'Por admin'}).status_code,200)
        self.assertEqual(self.post('/registros/clientes',cajero,{'nombre_razon_social':'Cliente de cajero'}).status_code,200)
        eventos=self.eventos(admin,'?modulo=clientes&accion=CREAR')['data']
        self.assertEqual(eventos[0]['id_usuario'],2)
        with self.engine.begin() as c:
            c.exec_driver_sql("INSERT INTO categoria(nombre) VALUES('Directa posterior')")
        evento=self.eventos(admin,'?modulo=categorias&accion=CREAR')['data'][0]
        self.assertIsNone(evento['id_usuario']);self.assertEqual(evento['origen'],'base_datos')

    def test_bitacora_conserva_perfil_que_autorizo_cambio_del_propio_rol(self):
        token=self.login()
        with self.engine.begin() as c:
            c.exec_driver_sql('UPDATE usuario SET id_rol=1 WHERE id_usuario=4')
        r=self.client.put('/api/registros/usuarios/1',headers=self.headers(token),json={'id_rol':2})
        self.assertEqual(r.status_code,200,r.json)
        with self.engine.connect() as c:
            evento=c.execute(text("SELECT id_usuario,usuario_rol,antes,despues FROM bitacora_auditoria WHERE accion='CAMBIAR_ROL' AND id_usuario=1 ORDER BY id_evento DESC LIMIT 1")).mappings().one()
        self.assertEqual(evento['usuario_rol'],'Administrador')
        self.assertEqual(evento['antes']['id_rol'],1);self.assertEqual(evento['despues']['id_rol'],2)


    def test_cinco_informes_y_exportaciones_reales(self):
        from openpyxl import load_workbook
        from io import BytesIO
        token=self.login()
        for tipo in ('ventas','productos','inventario','compras','caja'):
            r=self.client.get('/api/informes/'+tipo,headers=self.headers(token))
            self.assertEqual(r.status_code,200,r.json)
            for formato in ('pdf','excel'):
                r=self.client.get('/api/informes/'+tipo+'/'+formato,headers=self.headers(token))
                self.assertEqual(r.status_code,200)
                if formato=='pdf':self.assertTrue(r.data.startswith(b'%PDF'))
                else:
                    libro=load_workbook(BytesIO(r.data));self.assertTrue(libro.active.max_row>=8);libro.close()

    def test_informes_permisos_y_alcance_propio(self):
        token=self.login('cajero1')
        r=self.post('/ventas',token,self.payload());self.assertEqual(r.status_code,200,r.json)
        otro=self.login('cajero2')
        for tipo in ('ventas','productos','caja'):
            r=self.client.get('/api/informes/'+tipo,headers=self.headers(otro));self.assertEqual(r.status_code,200)
            self.assertEqual(r.json['data'],[])
        for tipo in ('compras','inventario'):
            for sufijo in ('','/pdf','/excel'):
                self.assertEqual(self.client.get('/api/informes/'+tipo+sufijo,headers=self.headers(token)).status_code,403)
        bodega=self.login('bodeguero')
        self.assertEqual(self.client.get('/api/informes/inventario',headers=self.headers(bodega)).status_code,200)
        self.assertEqual(self.client.get('/api/informes/ventas/pdf',headers=self.headers(bodega)).status_code,403)

    def test_informes_rangos_y_anulaciones(self):
        token=self.login()
        self.assertEqual(self.client.get('/api/informes/ventas/pdf?desde=2026-10-20&hasta=2026-10-01',headers=self.headers(token)).status_code,400)
        self.assertEqual(self.client.get('/api/informes/ninguno',headers=self.headers(token)).status_code,404)
        self.assertEqual(self.client.get('/api/informes/caja').status_code,401)


    def test_externos_autenticacion_validacion_y_caida_independiente_del_negocio(self):
        from unittest.mock import patch
        from backend import externos
        self.assertEqual(self.client.get('/api/externos/cambio').status_code,401)
        for usuario in ('admin','cajero1','bodeguero'):
            token=self.login(usuario)
            externos._cache.clear()
            with patch('backend.externos.obtener_json',side_effect=TimeoutError('No mostrar esto')):
                r=self.client.get('/api/externos/dominio?dominio=gmail.com',headers=self.headers(token))
            self.assertEqual(r.status_code,200);self.assertFalse(r.json['ok']);self.assertNotIn('No mostrar',r.text)
            self.assertEqual(self.client.get('/api/registros/productos',headers=self.headers(token)).status_code,200)
            self.assertEqual(self.client.get('/api/externos/feriados?anio=texto',headers=self.headers(token)).status_code,400)
            self.assertEqual(self.client.get('/api/externos/feriados?anio=1900',headers=self.headers(token)).status_code,400)
            self.assertEqual(self.client.get('/api/externos/desconocido',headers=self.headers(token)).status_code,404)


def main():
    from mfa_casos import MFATests
    from pagos_casos import PagosTests
    config=configuracion_origen()
    with BaseTemporalPostgres(config) as base:
        with base.engine.begin() as c:
            c.exec_driver_sql((ROOT/'sql/05_sesiones_token.sql').read_text(encoding='utf-8'))
        api_config={**config,'TESTING':True,'SECRET_KEY':'clave-de-pruebas-api-aisladas-de-al-menos-32-caracteres',
                    'SQLALCHEMY_DATABASE_URI':config['SQLALCHEMY_DATABASE_URI'].set(database=base.nombre)}
        APIJWTTests.aplicacion=create_app(api_config)
        for clase in (APIJWTTests,MFATests,PagosTests):
            clase.engine=base.engine;clase.app=base.modulo;clase.aplicacion=APIJWTTests.aplicacion
        suite=unittest.TestSuite(unittest.defaultTestLoader.loadTestsFromTestCase(clase) for clase in (APIJWTTests,MFATests,PagosTests))
        result=unittest.TextTestRunner(verbosity=2).run(suite)
        with APIJWTTests.aplicacion.app_context():db.session.remove();db.engine.dispose()
    reporte={'pruebas':result.testsRun,'errores':len(result.errors),'fallos':len(result.failures),
             'base_temporal_eliminada':not base.creada}
    (ROOT/'docs/resultado_api.json').write_text(json.dumps(reporte,indent=2)+'\n',encoding='utf-8')
    return 0 if result.wasSuccessful() and not base.creada else 1


if __name__=='__main__':raise SystemExit(main())
